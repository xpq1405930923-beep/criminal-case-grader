#!/usr/bin/env python3
"""Validate grading data and export an editable two-column DOCX; no legal inference."""
import argparse
from decimal import Decimal, InvalidOperation
import json
from pathlib import Path
import re
import sys

COLORS = {"correct": "006400", "error": "C00000", "neutral": "000000", "uncertain": "000000"}
BLUE = "0070C0"
LABELS = {"correct": "正确", "partial": "部分正确", "incorrect": "错误", "missing": "遗漏", "pending": "待核"}
ZERO = Decimal("0")


class DataError(ValueError):
    pass


def require(condition, message):
    if not condition:
        raise DataError(message)


def string(value, where, empty=False):
    require(isinstance(value, str), where + "必须是字符串")
    require(empty or bool(value.strip()), where + "不能为空")
    require(not re.search(r"[\x00-\x08\x0b-\x1f\ud800-\udfff\ufffe\uffff]", value),
            where + "含不可写入Word的字符或CR换行；转录确认前请统一为LF换行")
    return value


def number(value, where):
    require(not isinstance(value, bool) and isinstance(value, (int, float, Decimal)), where + "必须是数字")
    try:
        result = Decimal(str(value))
    except InvalidOperation:
        raise DataError(where + "不是有效数字")
    require(result.is_finite(), where + "不能是NaN或无穷大")
    require(result.as_tuple().exponent >= -2, where + "最多保留两位小数")
    return result


def sequence(value, where, nonempty=False):
    require(isinstance(value, list), where + "必须是数组")
    require(not nonempty or bool(value), where + "不能为空")
    return value


def obj(value, where):
    require(isinstance(value, dict), where + "必须是对象")
    return value


def validate(data):
    obj(data, "报告")
    string(data.get("title"), "title")
    string(data.get("grading_basis"), "grading_basis")
    if "student" in data:
        string(data["student"], "student", empty=True)
    for key in ("ocr_notes", "legal_sources"):
        for text in sequence(data.get(key, []), key):
            string(text, key)
    cases = sequence(data.get("cases"), "cases", True)
    case_ids = set()
    for case in cases:
        obj(case, "case")
        cid = string(case.get("id"), "案例id")
        require(cid not in case_ids, "案例id重复：" + cid)
        case_ids.add(cid)
        string(case.get("title"), cid + ".title")
        questions = sequence(case.get("questions"), cid + ".questions", True)
        lookup = {}
        for q in questions:
            obj(q, cid + "的小问")
            qid = string(q.get("id"), "小问id")
            require(qid not in lookup, cid + "的小问id重复：" + qid)
            lookup[qid] = q
            string(q.get("raw_text"), qid + ".raw_text", empty=True)
        case_max = ZERO
        point_ids = set()
        for q in questions:
            qid = q["id"]
            string(q.get("prompt"), qid + ".prompt")
            string(q.get("model_answer"), qid + ".model_answer")
            maximum = number(q.get("max_score"), qid + ".max_score")
            require(maximum > 0, qid + "满分必须大于0")
            case_max += maximum
            segments = sequence(q.get("segments"), qid + ".segments")
            for segment in segments:
                obj(segment, qid + "文字片段")
                string(segment.get("text"), qid + "片段text", empty=True)
                require(segment.get("kind") in COLORS, qid + "片段kind仅允许correct/error/neutral/uncertain；左栏不能补遗漏")
                require(isinstance(segment.get("struck", False), bool), qid + "片段struck必须为布尔值")
            require("".join(s["text"] for s in segments) == q["raw_text"], qid + "的片段拼接与原文不一致，禁止改写原文")
            for text in sequence(q.get("advice", []), qid + ".advice"):
                string(text, qid + ".advice")
            point_max = ZERO
            for p in sequence(q.get("points"), qid + ".points", True):
                obj(p, qid + "采分点")
                pid = string(p.get("id"), "采分点id")
                require(pid not in point_ids, cid + "采分点id重复：" + pid)
                point_ids.add(pid)
                string(p.get("criterion"), pid + ".criterion")
                string(p.get("comment"), pid + ".comment")
                cap = number(p.get("max_score"), pid + ".max_score")
                require(cap > 0, pid + "满分必须大于0")
                point_max += cap
                status = p.get("status")
                require(status in LABELS, pid + "状态无效")
                require("score" in p, pid + "缺少score")
                if status == "pending":
                    require(p["score"] is None, pid + "待核点score必须为null")
                else:
                    score = number(p["score"], pid + ".score")
                    require(0 <= score <= cap, pid + "得分超出0到满分的范围")
                    if status == "correct":
                        require(score == cap, pid + "正确点应得满分")
                    elif status == "partial":
                        require(0 < score < cap, pid + "部分正确应大于0且小于满分")
                    else:
                        require(score == 0, pid + "错误或遗漏点必须为0分")
                missing = string(p.get("missing", ""), pid + ".missing", empty=True)
                if status == "missing":
                    require(bool(missing.strip()), pid + "遗漏点必须写明missing")
                if status in ("correct", "pending"):
                    require(not missing, pid + "正确或待核点不能同时标为确定遗漏")
                evidence = sequence(p.get("evidence", []), pid + ".evidence")
                if status in ("correct", "partial", "incorrect"):
                    require(bool(evidence), pid + "必须提供考生原文证据")
                if status == "missing":
                    require(not evidence, pid + "纯遗漏点不能伪造原文证据")
                for ev in evidence:
                    obj(ev, pid + "证据")
                    target = string(ev.get("question_id"), pid + "证据question_id")
                    quote = string(ev.get("quote"), pid + "证据quote")
                    require(target in lookup, pid + "证据引用不存在的小问")
                    require(quote in lookup[target]["raw_text"], pid + "引文不是所引用小问原文的连续片段")
            require(point_max == maximum, qid + "采分点满分之和不等于该问满分")
        require(case_max == 20, cid + "全部小问满分之和必须为20分")
    return data


def fmt(n):
    return format(Decimal(n).normalize(), "f")


def tally(points):
    earned = sum((Decimal(str(p["score"])) for p in points if p["score"] is not None), ZERO)
    pending = sum((Decimal(str(p["max_score"])) for p in points if p["score"] is None), ZERO)
    maximum = sum((Decimal(str(p["max_score"])) for p in points), ZERO)
    return earned, pending, maximum


def score_text(points):
    earned, pending, maximum = tally(points)
    if pending:
        return f"暂得{fmt(earned)}分；待核{fmt(pending)}分；区间{fmt(earned)}—{fmt(earned + pending)}分（满分{fmt(maximum)}分）"
    return f"得分{fmt(earned)} / {fmt(maximum)}分"


def export_docx(data, path, font=None):
    validate(data)
    try:
        from docx import Document
        from docx.shared import Cm, Pt, RGBColor
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    except ImportError:
        raise DataError("缺少python-docx；请使用Codex配套Python，或在所选环境安装python-docx")

    font = font or ("Songti SC" if sys.platform == "darwin" else "SimSun" if sys.platform == "win32" else "Noto Serif CJK SC")

    def set_font(rpr):
        fonts = rpr.get_or_add_rFonts()
        for key in list(fonts.attrib):
            if key.endswith("Theme"):
                del fonts.attrib[key]
        for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn("w:" + slot), font)

    def run(paragraph, text, color="000000", bold=False, struck=False):
        r = paragraph.add_run(text)
        set_font(r._element.get_or_add_rPr())
        r.font.color.rgb = RGBColor.from_string(color)
        r.bold = bold
        r.font.strike = struck
        return r

    def para(parent, text="", bold=False, color="000000"):
        p = parent.add_paragraph()
        run(p, text, color, bold)
        return p

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.7)
    section.left_margin = section.right_margin = Cm(1.6)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        set_font(style._element.get_or_add_rPr())
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(5)
    # Some installed base templates add a blue title rule; do not inherit it.
    for border in doc.styles.element.findall(".//" + qn("w:pBdr")):
        border.getparent().remove(border)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.18
    doc.styles["Title"].font.size = Pt(20)
    doc.styles["Heading 1"].font.size = Pt(14)
    doc.styles["Heading 2"].font.size = Pt(12)
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.title = data["title"]
    doc.add_paragraph(data["title"], "Title")
    if data.get("student"):
        para(doc, "考生：" + data["student"])
    all_points = [p for c in data["cases"] for q in c["questions"] for p in q["points"]]
    para(doc, "总评  " + score_text(all_points), bold=True)
    para(doc, "评分依据：" + data["grading_basis"])
    for note in data.get("ocr_notes", []):
        para(doc, "识别复核：" + note)
    legend = doc.add_paragraph()
    run(legend, "标注说明：")
    run(legend, "深绿为得分内容", "006400")
    run(legend, "；")
    run(legend, "红色为错误", "C00000")
    # Blue is confined to right-hand grading cells, including the legend explanation.
    run(legend, "；蓝色遗漏仅在右栏；黑色为其他原文或待核内容。")
    for case in data["cases"]:
        doc.add_heading(case["title"], level=1)
        points = [p for q in case["questions"] for p in q["points"]]
        para(doc, score_text(points), bold=True)
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        widths = (Cm(8.3), Cm(9.5))
        for col, width in zip(table.columns, widths):
            col.width = width
        props = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            b = OxmlElement("w:" + edge)
            for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                b.set(qn("w:" + key), value)
            borders.append(b)
        props.append(borders)
        margins = OxmlElement("w:tblCellMar")
        for edge in ("top", "left", "bottom", "right"):
            m = OxmlElement("w:" + edge)
            m.set(qn("w:w"), "100")
            m.set(qn("w:type"), "dxa")
            margins.append(m)
        props.append(margins)
        header = table.rows[0]
        repeat = OxmlElement("w:tblHeader")
        header._tr.get_or_add_trPr().append(repeat)
        for cell, width, title in zip(header.cells, widths, ("考生原本答案", "这一问的采分点情况")):
            cell.width = width
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "E8EDF2")
            cell._tc.get_or_add_tcPr().append(shading)
            run(cell.paragraphs[0], title, bold=True)
            cell.paragraphs[0].paragraph_format.keep_with_next = True
        for q in case["questions"]:
            row = table.add_row()
            left, right = row.cells
            for cell, width in zip(row.cells, widths):
                cell.width = width
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            # One paragraph with explicit line breaks preserves raw_text byte-for-byte as text.
            original = left.paragraphs[0]
            for segment in q["segments"]:
                run(original, segment["text"], COLORS[segment["kind"]], struck=segment.get("struck", False))
            run(right.paragraphs[0], q["id"] + "  " + q["prompt"], bold=True)
            right.paragraphs[0].paragraph_format.keep_with_next = True
            for p in q["points"]:
                score = "待核" if p["score"] is None else fmt(Decimal(str(p["score"])))
                color = "006400" if p["status"] == "correct" else "C00000" if p["status"] == "incorrect" else "000000"
                paragraph = para(right, f'{p["id"]}  {p["criterion"]}  {score}/{fmt(Decimal(str(p["max_score"])))}分  {LABELS[p["status"]]}', bold=True, color=color)
                paragraph.paragraph_format.keep_with_next = True
                for ev in p.get("evidence", []):
                    para(right, f'原文（{ev["question_id"]}）：“{ev["quote"]}”')
                para(right, p["comment"])
                if p.get("missing"):
                    lost = Decimal(str(p["max_score"])) - Decimal(str(p["score"]))
                    para(right, f'遗漏：{p["missing"]}（本点未得{fmt(lost)}分）', color=BLUE)
            para(right, "本问小计  " + score_text(q["points"]), bold=True)
        doc.add_heading("简明示范答案", level=2)
        for q in case["questions"]:
            p = para(doc, q["id"] + "  " + q["prompt"], bold=True)
            p.paragraph_format.keep_with_next = True
            for line in q["model_answer"].split("\n"):
                para(doc, line)
        advice = [(q["id"], text) for q in case["questions"] for text in q.get("advice", [])]
        if advice:
            doc.add_heading("改进建议", level=2)
            for qid, text in advice:
                para(doc, qid + "：" + text)
    if data.get("legal_sources"):
        doc.add_heading("核验依据", level=1)
        for source in data["legal_sources"]:
            para(doc, source)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path, nargs="?")
    parser.add_argument("--validate-only", action="store_true")
    parser.add_argument("--font", help="已安装的中文字体名；默认macOS宋体Songti SC、Windows SimSun、Linux Noto Serif CJK SC")
    args = parser.parse_args(argv)
    try:
        with args.input.open(encoding="utf-8-sig") as f:
            data = json.load(f, parse_float=Decimal)
        validate(data)
        if args.validate_only:
            print("校验通过：分值、原文及证据映射一致。")
            return 0
        require(args.output is not None, "请指定输出.docx，或使用--validate-only")
        require(args.output.suffix.lower() == ".docx", "输出必须为.docx")
        require(args.input.resolve() != args.output.resolve(), "输出不能覆盖评分JSON")
        export_docx(data, args.output, args.font)
        print("已生成Word：" + str(args.output.resolve()))
        return 0
    except (DataError, ValueError, OSError) as exc:
        print("生成失败：" + str(exc), file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
