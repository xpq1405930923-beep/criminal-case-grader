#!/usr/bin/env python3
"""Offline invariants for the Word exporter and OCR client. No user data or network."""
from copy import deepcopy
from decimal import Decimal
import contextlib
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile
import xml.etree.ElementTree as ET

import build_report as report
import ocr_aistudio as ocr

FIXTURE = Path(__file__).resolve().parents[1] / "assets" / "report-example.json"


class ReportTests(unittest.TestCase):
    def setUp(self):
        self.data = json.loads(FIXTURE.read_text(encoding="utf-8"))
        self.q1, self.q2 = self.data["cases"][0]["questions"]

    def test_totals_and_pending(self):
        report.validate(self.data)
        points = self.q1["points"] + self.q2["points"]
        self.assertEqual(report.tally(points), (Decimal(6), Decimal(4), Decimal(20)))
        self.assertIn("6—10", report.score_text(points))

    def test_original_change_rejected(self):
        self.q1["segments"][1]["text"] += "擅自改写"
        with self.assertRaises(report.DataError):
            report.validate(self.data)

    def test_two_column_docx_roundtrip_and_colors(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "test.docx"
            report.export_docx(self.data, path)
            doc = Document(path)
            self.assertEqual(len(doc.tables), 1)
            table = doc.tables[0]
            self.assertEqual(len(table.columns), 2)
            for row, q in zip(table.rows[1:], (self.q1, self.q2)):
                self.assertEqual(row.cells[0].text, q["raw_text"])
                colors = {str(r.font.color.rgb) for p in row.cells[0].paragraphs for r in p.runs}
                self.assertNotIn(report.BLUE, colors)
            colors = {str(r.font.color.rgb) for p in table.rows[1].cells[0].paragraphs for r in p.runs}
            self.assertTrue({"006400", "C00000"} <= colors)
            right = table.rows[1].cells[1]
            self.assertTrue(any(str(r.font.color.rgb) == report.BLUE for p in right.paragraphs for r in p.runs))
            with zipfile.ZipFile(path) as z:
                root = ET.fromstring(z.read("word/document.xml"))
                ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                self.assertEqual(len(root.findall(".//w:tblHeader", ns)), 1)

    def test_cross_question_evidence(self):
        self.q1["points"][0]["evidence"] = [{"question_id": "Q2", "quote": "2. "}]
        report.validate(self.data)
        self.q1["points"][0]["evidence"][0]["quote"] = "不存在的内容"
        with self.assertRaises(report.DataError):
            report.validate(self.data)

    def test_invalid_scores_and_case_total(self):
        for bad in (True, float("nan"), float("inf"), -1, 5, 1.123):
            trial = deepcopy(self.data)
            trial["cases"][0]["questions"][0]["points"][0]["score"] = bad
            with self.subTest(bad=bad), self.assertRaises(report.DataError):
                report.validate(trial)
        self.q1["max_score"] = 13
        with self.assertRaises(report.DataError):
            report.validate(self.data)

    def test_blank_original_is_not_comment(self):
        self.q2["raw_text"] = ""
        self.q2["segments"] = []
        self.q2["points"][1]["evidence"] = []
        report.validate(self.data)

    def test_multiline_tabs_strike_preserved(self):
        from docx import Document
        self.q1["raw_text"] += "\n\n\t划除的文字"
        self.q1["segments"].append({"text": "\n\n\t划除的文字", "kind": "neutral", "struck": True})
        with tempfile.TemporaryDirectory() as folder:
            path = report.export_docx(self.data, Path(folder) / "multiline.docx")
            left = Document(path).tables[0].rows[1].cells[0]
            self.assertEqual(left.text, self.q1["raw_text"])
            self.assertTrue(left.paragraphs[0].runs[-1].font.strike)


class OCRTests(unittest.TestCase):
    def test_extract_ocr_and_layout(self):
        records = [{"result": {"ocrResults": [{"prunedResult": {"rec_texts": ["甲", "乙"], "rec_scores": [0.9, 0.6]}}]}}]
        self.assertEqual(ocr.extract_pages(records)[0]["text"], "甲\n乙")
        layout = [{"result": {"layoutParsingResults": [{"markdown": {"text": "文档文字"}}]}}]
        self.assertEqual(ocr.extract_pages(layout)[0]["text"], "文档文字")
        with self.assertRaises(ocr.OCRError):
            ocr.extract_pages([{"result": {}}])

    def test_polling_state_machine(self):
        states = iter([{"state": "pending"}, {"state": "running"}, {"state": "done", "resultUrl": {"jsonUrl": "https://result.bcebos.com/out"}}])
        clock = [0.0]
        result = ocr.poll_job("ocrjob-test", "dummy", 2, 5, 1,
                              api=lambda *a: next(states), now=lambda: clock[0],
                              sleep=lambda n: clock.__setitem__(0, clock[0] + n))
        self.assertEqual(result["state"], "done")
        for state in ("failed", "unknown", "done"):
            with self.subTest(state=state), self.assertRaises(ocr.OCRError):
                ocr.poll_job("job1", "dummy", 1, 2, 1, api=lambda *a: {"state": state})

    def test_no_token_forwarding_to_other_hosts(self):
        with self.assertRaises(ocr.OCRError):
            ocr.api_request("https://example.com/api/v2/ocr/jobs", "dummy", 1)
        with self.assertRaises(ocr.OCRError):
            ocr.download_results("https://example.com/result", 1)
        with patch.object(ocr, "request_bytes", return_value=b'{"result": {}}') as request:
            ocr.download_results("https://result.bcebos.com/out.jsonl", 1)
            self.assertNotIn("headers", request.call_args.kwargs)

    def test_one_post_checkpoint_and_redaction(self):
        with tempfile.TemporaryDirectory() as folder:
            source, output = Path(folder) / "test.png", Path(folder) / "ocr.json"
            source.write_bytes(b"synthetic-input")
            with patch.dict(ocr.os.environ, {"AISTUDIO_ACCESS_TOKEN": "dummy-private-token"}), \
                 patch.object(ocr, "api_request", return_value={"jobId": "ocrjob-test"}) as submit, \
                 patch.object(ocr, "poll_job", side_effect=ocr.OCRError("poll_timeout", "dummy-private-token")), \
                 contextlib.redirect_stderr(io.StringIO()):
                self.assertEqual(ocr.main([str(source), "--output", str(output)]), 3)
                self.assertEqual(submit.call_count, 1)
            contents = output.read_text(encoding="utf-8")
            self.assertIn("ocrjob-test", contents)
            self.assertNotIn("dummy-private-token", contents)

    def test_timeout_bounded(self):
        clock = [0.0]
        with self.assertRaises(ocr.OCRError) as error:
            ocr.poll_job("job1", "dummy", 1, 2, 1, api=lambda *a: {"state": "pending"},
                          now=lambda: clock[0], sleep=lambda n: clock.__setitem__(0, clock[0] + n))
        self.assertEqual(error.exception.kind, "poll_timeout")
        self.assertEqual(clock[0], 2)


if __name__ == "__main__":
    unittest.main(verbosity=2)
